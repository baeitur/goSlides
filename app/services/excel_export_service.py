"""Excel export for participant/registrant lists."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_registrants_excel(activity, registrants):
    """Generate Excel buffer with participant list for the activity."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"
    
    # Define styles
    header_fill = PatternFill(start_color="1BA3A8", end_color="1BA3A8", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    
    data_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # Add title and activity info
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "Go Slides – Participant List"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("A2:H2")
    activity_cell = ws["A2"]
    activity_cell.value = activity.title
    activity_cell.font = Font(bold=True, size=12)
    activity_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("A3:H3")
    date_cell = ws["A3"]
    date_cell.value = f"Activity date: {activity.date.strftime('%d %B %Y') if activity.date else 'TBA'}"
    date_cell.font = Font(size=10)
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Add empty row for spacing
    ws.append([])
    
    # Add headers
    headers = ["#", "Name", "School", "Email", "Phone", "Status", "Attended", "Check-in Code"]
    header_row = 5
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Set column widths
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 20
    
    # Add data rows
    for i, r in enumerate(registrants, 1):
        row = 5 + i
        ws.cell(row=row, column=1).value = i
        ws.cell(row=row, column=2).value = r.name or ""
        ws.cell(row=row, column=3).value = r.school or ""
        ws.cell(row=row, column=4).value = r.email or ""
        ws.cell(row=row, column=5).value = r.phone or ""
        ws.cell(row=row, column=6).value = r.status or ""
        ws.cell(row=row, column=7).value = "Yes" if r.attended_at else "No"
        ws.cell(row=row, column=8).value = r.check_in_code or ""
        
        # Apply formatting to data rows
        for col_num in range(1, 9):
            cell = ws.cell(row=row, column=col_num)
            cell.alignment = data_alignment
            cell.border = border
            # Alternate row colors
            if i % 2 == 0:
                cell.fill = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
    
    # Create buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
